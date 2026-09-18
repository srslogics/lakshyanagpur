from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
import json

import pytest
from openpyxl import load_workbook

from app.models import (
    Batch, BiometricAttendanceDay, BiometricImportBatch, ClassSession,
    DeviceAttendanceIdentity, Examination, ExaminationParticipant, ExaminationResult,
    FeeAgreement, FeeInstallment, Guardian, InventoryItem, InventoryMovement, Lead,
    LeadActivity, PaymentTransaction, Room, StaffAttendanceWorkday, StaffPayroll,
    Student, StudentGuardian, Subject, User, UserModulePermission,
    Assignment, AssignmentRecipient, AssignmentDownload, AssignmentMaterial,
    StudentAcademicProfile, StudentSubjectSelection, Enrollment, Notice, FacultyTeachingAssignment,
)
from app.report_catalog import REPORTS
from app.payroll import calculate_payroll
from app.routers.payroll import _people
from app.security import create_token
from test_report_exports import _finance_records, _workbook
from test_payroll import _payroll_people


@pytest.mark.parametrize("definition", REPORTS, ids=lambda r: r["id"])
def test_every_catalog_entry_downloads_all_declared_sheets(client, owner_headers, definition):
    workbook = _workbook(client, owner_headers, definition["id"])
    assert workbook.sheetnames == definition["sheets"]
    for sheet in workbook:
        assert sheet["A6"].value
        assert sheet.auto_filter.ref.startswith("A6:")
        assert sheet["A7"].value == "No records found."


def seed_details(database):
    _finance_records(database)
    owner, staff, director, accounts = _payroll_people(database)
    student = database.query(Student).one()
    agreement = database.query(FeeAgreement).one()
    guardian = Guardian(full_name="Sample Guardian", mobile="09000000000", verified=True)
    batch, subject, room = Batch(name="Essential", program="JEE"), Subject(name="Physics", code="QA-PHY", program="JEE"), Room(name="Room A")
    database.add_all([guardian, batch, subject, room])
    database.flush()
    database.add(StudentGuardian(student_id=student.id, guardian_id=guardian.id, relationship="father"))
    database.add(FeeInstallment(student_id=student.id, fee_agreement_id=agreement.id, due_date=date(2026, 8, 10), amount=10000, created_by=owner.id, notes="First instalment"))
    database.add(PaymentTransaction(student_id=student.id, fee_agreement_id=agreement.id, transaction_date=date(2026, 8, 10), amount=500, method="cash", transaction_type="refund", status="posted", reconciliation_status="ready", source_note="Refund"))
    database.add(ClassSession(batch_id=batch.id, subject_id=subject.id, faculty_id=staff.id, room_id=room.id, starts_at=datetime(2026, 8, 10, 4, tzinfo=timezone.utc), ends_at=datetime(2026, 8, 10, 5, tzinfo=timezone.utc), status="cancelled", notes="Rescheduled"))
    exam = Examination(name="Unit test", batch_id=batch.id, subject_id=subject.id, faculty_id=staff.id, scheduled_at=datetime(2026, 8, 10, 4, tzinfo=timezone.utc), max_marks=50, pass_marks=17.5, created_by=owner.id, status="draft")
    second = Student(full_name="Pending Student", admission_number="QA-002", status="active")
    lead = Lead(student="Enquiry Sample", mobile="09000000001", program="JEE", parent="Guardian", source="walk_in", counsellor="Desk", stage="New", next_action="Call", created_at=datetime(2026, 8, 1, tzinfo=timezone.utc))
    item = InventoryItem(sku="QA-001", name="Physics book", category="book", quantity_on_hand=9, created_by=owner.id)
    database.add_all([exam, second, lead, item])
    database.flush()
    for s in (student, second):
        database.add(ExaminationParticipant(exam_id=exam.id, student_id=s.id, admission_number=s.admission_number, full_name=s.full_name))
    database.add(ExaminationResult(exam_id=exam.id, student_id=student.id, marks_obtained=0, result_status="graded", entered_by=owner.id))
    database.add(LeadActivity(lead_id=lead.id, kind="call", note="Follow up after requested creation range", actor_id=owner.id, created_at=datetime(2026, 9, 1, tzinfo=timezone.utc)))
    database.add(InventoryMovement(item_id=item.id, movement_type="issue", quantity_delta=-1, balance_after=9, occurred_on=date(2026, 8, 10), student_id=student.id, reason="Issued book", created_by=owner.id))
    database.query(Enrollment).filter_by(student_id=student.id).one().batch = batch.name
    database.add(StudentAcademicProfile(student_id=student.id, source_student_code="E-01", batch_name=batch.name, mentor_name="Sample Mentor"))
    database.add(StudentSubjectSelection(student_id=student.id, subject_name="Physics", source_value="Yes"))
    database.add(FacultyTeachingAssignment(faculty_id=staff.id, batch_id=batch.id, subject_id=subject.id, created_by=owner.id))
    assignment = Assignment(title="Physics practice", batch_id=batch.id, subject_id=subject.id, instructions="Complete exercise 1", due_at=datetime(2026, 8, 10, 18, 30, tzinfo=timezone.utc), external_url="", created_by=owner.id)
    database.add(assignment)
    database.flush()
    database.add(AssignmentRecipient(assignment_id=assignment.id, student_id=student.id, status="completed"))
    database.add(AssignmentDownload(assignment_id=assignment.id, student_id=second.id, download_count=2))
    database.add(AssignmentMaterial(assignment_id=assignment.id, filename="exercise.pdf", size_bytes=8, content=b"%PDF-QA", expires_at=datetime(2026, 8, 12, tzinfo=timezone.utc)))
    database.add(Notice(title="Test notice", body="=1+1", audience="students", batch_id=batch.id, subject_id=subject.id, created_by=owner.id, created_at=datetime(2026, 8, 10, 18, 30, tzinfo=timezone.utc)))
    imported = database.query(BiometricImportBatch).one()
    database.add(StaffAttendanceWorkday(import_batch_id=imported.id, device_key="payroll-device", device_user_id="41", staff_user_id=staff.id, attendance_date=date(2026, 8, 1), attendance_status="half_day", work_duration_minutes=255, overtime_minutes=15, punch_count=4))
    database.add(BiometricAttendanceDay(import_batch_id=imported.id, device_key="payroll-device", device_user_id="42", staff_user_id=director.id, attendance_date=date(2026, 8, 1), first_punch_at=datetime(2026, 8, 1, 5, tzinfo=timezone.utc)))
    database.flush()
    people = _people(database, "2026-08")
    key, person = next(iter(people.items()))
    calc = calculate_payroll("2026-08", Decimal("31000.50"), Decimal("1.5"), Decimal("1000.25"))
    database.add(StaffPayroll(person_key=key, month="2026-08", monthly_salary=31000.50, advance_given=1000.25, absent_days=1.5, attendance_fingerprint=person["attendanceFingerprint"], status="finalized", updated_by=owner.id, snapshot={"fullName": person["fullName"], "attendance": person, "calculation": calc}))
    database.commit()
    return student, staff, calc


def test_populated_detailed_reports_and_saved_files(client, database, owner_headers, tmp_path):
    student, staff, calc = seed_details(database)
    books = {}
    for definition in REPORTS:
        name = definition["id"]
        response = client.get(f"/api/reports/export/{name}?format=xlsx" + ("&month=2026-08" if name == "payroll" else ""), headers=owner_headers)
        assert response.status_code == 200, response.text
        (tmp_path / f"{name}.xlsx").write_bytes(response.content)
        books[name] = load_workbook(BytesIO(response.content))
    (tmp_path / "catalog.json").write_text(json.dumps(REPORTS))
    assert books["students"]["Guardian contacts"]["D7"].value == "09000000000"
    assert books["fees"]["Agreement details"]["L7"].value
    ledger = books["payments"].active
    assert ledger["G8"].value == -500
    assert books["installments"].active["D7"].value == 10000
    assert books["inventory"]["Stock movements"]["E7"].value == -1
    assert books["timetable"].active["G7"].value == "cancelled"
    assert books["timetable"]["Faculty allocations"]["A7"].value == staff.full_name
    assert books["academics"]["Academic profiles"]["G7"].value == "Sample Mentor"
    assert books["academics"]["Subject selections"]["D7"].value == "Physics"
    assert books["assignments"]["Assignments"]["J7"].value == 1
    assert books["assignments"]["Assignments"]["K7"].value == 1
    assert books["assignments"]["Student progress"]["F7"].value == "completed"
    assert books["assignments"]["Student progress"]["E8"].value == "No"
    assert books["assignments"]["Student progress"]["J8"].value == 2
    assert books["assignments"]["Material details"]["C7"].value == "exercise.pdf"
    assert books["announcements"].active["C7"].value == "=1+1"
    assert books["announcements"].active["C7"].data_type == "s"
    exam = books["examinations"]["Student results"]
    statuses = {exam.cell(row, 4).value: (exam.cell(row, 9).value, exam.cell(row, 10).value) for row in (7, 8)}
    assert statuses[student.full_name] == (0, "graded")
    assert statuses["Pending Student"] == (None, "not_entered")
    assert books["examinations"]["Examinations"]["I7"].value == 17.5
    staff_sheet = books["staff-attendance"]["Staff daily"]
    assert staff_sheet.max_row == 8  # The monthly row replaces the same day's punches.
    assert staff_sheet["I8"].value == timedelta(minutes=255)
    assert staff_sheet["I8"].number_format == "[h]:mm"
    assert books["staff-attendance"]["Director daily"]["B7"].value == "Institute Director"
    totals = books["staff-attendance"]["Monthly totals"]
    director_row = next(row for row in totals.iter_rows(min_row=7, values_only=True) if row[1] == "Institute Director")
    assert director_row[8] is None
    assert director_row[11] == 1
    payroll = books["payroll"]["Payroll register"]
    assert payroll["F7"].value == 1.5
    assert payroll["K7"].value == float(calc["netPayable"])
    assert books["payroll"]["Saved attendance"]["E7"].value == timedelta(minutes=255)
    saved_attendance = books["payroll"]["Saved attendance"]
    assert saved_attendance["C8"].value == datetime(2026, 8, 15)
    assert saved_attendance["D8"].value == "present"
    assert saved_attendance["E8"].value is None
    assert saved_attendance["F8"].value is None
    # Finalized exports never recalculate a saved financial snapshot.
    database.query(StaffPayroll).one().monthly_salary = 99999
    database.commit()
    assert _workbook(client, owner_headers, "payroll", "&month=2026-08").active["K7"].value == float(calc["netPayable"])
    leads = _workbook(client, owner_headers, "admissions", "&from=2026-08-01&to=2026-08-01")
    assert leads["Follow-up history"]["C7"].value == datetime(2026, 9, 1, 5, 30)
    assert _workbook(client, owner_headers, "examinations", "&from=2026-08-11").active["A7"].value == "No records found."
    assert _workbook(client, owner_headers, "assignments", "&from=2026-08-10&to=2026-08-10").active["A7"].value == "No records found."
    assert _workbook(client, owner_headers, "assignments", "&from=2026-08-11&to=2026-08-11").active["B7"].value == "Physics practice"
    csv = client.get("/api/reports/export/announcements?format=csv&from=2026-08-11&to=2026-08-11", headers=owner_headers)
    assert "2026-08-11 00:00:00" in csv.text
    assert "'=1+1" in csv.text


def test_exports_enforce_source_permissions_and_custom_allowlists(client, database):
    user = User(full_name="Academic", mobile="9999999999", role="academic_coordinator", password_hash="unused")
    database.add(user)
    database.commit()
    headers = {"Authorization": f"Bearer {create_token(user)}"}
    ids = {r["id"] for r in client.get("/api/reports/exports", headers=headers).json()}
    assert "examinations" in ids
    assert not {"fees", "payments", "payroll", "inventory"} & ids
    assert {"academics", "assignments", "announcements"} <= ids
    assert "recordedPayments" not in client.get("/api/reports/overview", headers=headers).json()["metrics"]
    for name in ("fees", "payments", "payroll", "inventory"):
        assert client.get(f"/api/reports/export/{name}?format=xlsx", headers=headers).status_code == 403
    database.add(UserModulePermission(user_id=user.id, module="reports", can_read=True, can_create=False, can_edit=False))
    database.commit()
    assert {r["id"] for r in client.get("/api/reports/exports", headers=headers).json()} == {"audit"}
    assert client.get("/api/reports/export/examinations?format=xlsx", headers=headers).status_code == 403
    overview = client.get("/api/reports/overview", headers=headers).json()
    assert overview["leadFunnel"] == []
    assert overview["attendance"] == []
    assert set(overview["metrics"]) == {"activeUsers"}


def test_invalid_filters_never_silently_drop_sheets(client, owner_headers):
    for path in ("payroll?month=2026-13", "payroll?from=2026-01-01", "students?month=2026-08", "examinations?format=csv", "staff-attendance?from=2026-09-02&to=2026-09-01"):
        assert client.get(f"/api/reports/export/{path}", headers=owner_headers).status_code == 422


def test_staff_export_is_uncapped_and_excludes_test_staff(client, database, owner_headers):
    _, staff, _, _ = _payroll_people(database)
    imported = database.query(BiometricImportBatch).one()
    for index in range(510):
        day = date(2024, 1, 1) + timedelta(days=index)
        database.add(BiometricAttendanceDay(import_batch_id=imported.id, device_key="payroll-device", device_user_id="41", staff_user_id=staff.id, attendance_date=day, first_punch_at=datetime.combine(day, datetime.min.time())))
    database.commit()
    assert _workbook(client, owner_headers, "staff-attendance")["Staff daily"].max_row == 518
    staff.is_test_account = True
    database.commit()
    assert _workbook(client, owner_headers, "staff-attendance")["Staff daily"]["A7"].value == "No records found."
