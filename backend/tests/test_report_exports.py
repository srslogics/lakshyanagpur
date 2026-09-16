from datetime import date, datetime, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import (
    AttendanceEntry,
    AttendancePeriodSummary,
    AttendanceRegister,
    AuditLog,
    Batch,
    Subject,
    Room,
    ClassSession,
    Enrollment,
    FeeAgreement,
    PaymentTransaction,
    Student,
    User,
)
from app.security import create_token, hash_password


def _workbook(client, owner_headers, report, suffix=""):
    response = client.get(f"/api/reports/export/{report}?format=xlsx{suffix}", headers=owner_headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert response.headers["content-disposition"].endswith('.xlsx"')
    assert response.headers["cache-control"] == "no-store"
    return load_workbook(BytesIO(response.content))


@pytest.mark.parametrize("report,sheets", [
    ("students", ["Student register"]),
    ("fees", ["Open accounts", "Closed accounts"]),
    ("attendance", ["Daily attendance", "Class attendance", "Period summaries"]),
    ("audit", ["Audit trail"]),
])
def test_all_excel_downloads_open_with_headers_filters_and_print_layout(client, owner_headers, report, sheets):
    workbook = _workbook(client, owner_headers, report)
    assert workbook.sheetnames == sheets
    for sheet in workbook:
        assert sheet.freeze_panes == "C7"
        assert sheet.auto_filter.ref.startswith("A6:")
        assert sheet["A6"].font.color.rgb == "00FFFFFF"
        assert sheet.sheet_view.showGridLines is False
        assert sheet.page_setup.fitToWidth == 1
        assert sheet.page_setup.fitToHeight == 0
        assert sheet.print_title_rows == "$1:$6"
        assert sheet["A7"].value == "No records found."


def test_student_excel_preserves_identifiers_dates_and_literal_text(client, database, owner_headers):
    _finance_records(database)
    student = database.query(Student).one()
    student.mobile = "09000000991"
    student.full_name = '=HYPERLINK("https://invalid.example")'
    database.add(Student(admission_number="TEST-HIDDEN", full_name="Hidden", status="active", is_test_account=True))
    database.commit()
    sheet = _workbook(client, owner_headers, "students").active
    assert sheet["B7"].value == student.full_name
    assert sheet["B7"].data_type == "s"
    assert sheet["C7"].value == "09000000991"
    assert sheet["C7"].number_format == "@"
    assert sheet["H7"].value == datetime(2026, 7, 1)
    assert sheet.max_row == 7


def test_excel_fees_separate_closed_accounts_preserve_negative_receipts(client, database, owner_headers):
    _finance_records(database)
    agreement = database.query(FeeAgreement).one()
    database.add(PaymentTransaction(student_id=agreement.student_id, fee_agreement_id=agreement.id,
        transaction_date=date(2026, 7, 3), amount=6000, method="upi", transaction_type="refund",
        status="posted", reconciliation_status="ready", source_note="Test refund"))
    database.commit()
    sheet = _workbook(client, owner_headers, "fees")["Open accounts"]
    assert sheet["C7"].value == 80000
    assert sheet["D7"].value == -1000
    assert sheet["D7"].data_type == "n"
    assert sheet["E7"].value == 81000
    assert sheet["E9"].value == 81000
    database.query(Student).one().status = "forfeited"
    database.commit()
    workbook = _workbook(client, owner_headers, "fees")
    assert workbook["Open accounts"]["A7"].value == "No records found."
    assert workbook["Closed accounts"]["D7"].value == -1000
    assert workbook["Closed accounts"]["E7"].value == 0


def test_audit_dates_use_ist_and_exclude_test_accounts(client, database, owner_headers):
    owner = database.query(User).filter_by(role="owner").one()
    test_user = User(mobile="9000000899", full_name="Hidden test actor", role="owner", password_hash="unused", is_test_account=True)
    database.add(test_user)
    database.flush()
    for actor, at, action in [
        (owner, datetime(2026, 9, 15, 18, 30, tzinfo=timezone.utc), "included"),
        (owner, datetime(2026, 9, 15, 18, 29, tzinfo=timezone.utc), "before"),
        (owner, datetime(2026, 9, 16, 18, 30, tzinfo=timezone.utc), "after"),
        (test_user, datetime(2026, 9, 16, 10, tzinfo=timezone.utc), "test-hidden"),
    ]:
        database.add(AuditLog(actor_id=actor.id, created_at=at, action=action, entity_type="student", entity_id="sample"))
    database.commit()
    sheet = _workbook(client, owner_headers, "audit", "&from=2026-09-16&to=2026-09-16").active
    assert sheet.max_row == 7
    assert sheet["A7"].value == datetime(2026, 9, 16)
    assert sheet["D7"].value == "included"


def test_attendance_excel_uses_manual_priority_and_excludes_drafts_and_tests(client, database, owner_headers):
    _finance_records(database)
    owner = database.query(User).filter_by(role="owner").one()
    student = database.query(Student).one()
    test_student = Student(admission_number="TEST-ATT", full_name="Hidden", status="active", is_test_account=True)
    database.add(test_student)
    database.flush()
    for kind, status, day in [("manual", "submitted", 16), ("biometric", "submitted", 16), ("manual", "draft", 17)]:
        register = AttendanceRegister(register_kind=kind, attendance_date=date(2026, 9, day), batch_name="Tatva", stream_name="__all__", subject_name=kind, status=status)
        database.add(register)
        database.flush()
        for person in (student, test_student):
            database.add(AttendanceEntry(register_id=register.id, student_id=person.id, status="present" if kind == "manual" else "absent", reason="source", marked_by=owner.id))
    database.commit()
    sheet = _workbook(client, owner_headers, "attendance")["Daily attendance"]
    assert sheet.max_row == 7
    assert sheet["A7"].value == datetime(2026, 9, 16)
    assert sheet["E7"].value == "All programs"
    assert sheet["G7"].value == "present"
    assert sheet["I7"].value == "manual"
    empty = _workbook(client, owner_headers, "attendance", "&from=2026-09-17")["Daily attendance"]
    assert empty["A7"].value == "No records found."


def test_report_filters_fail_instead_of_silently_being_ignored(client, owner_headers):
    for report in ("students", "fees"):
        assert client.get(f"/api/reports/export/{report}?from=2026-01-01", headers=owner_headers).status_code == 422
    assert client.get("/api/reports/export/students?format=pdf", headers=owner_headers).status_code == 422


def test_csv_formula_safety_preserves_numeric_negative_values():
    from app.routers.reports import _safe_csv_cell
    assert _safe_csv_cell(" \t=1+1").startswith("'")
    assert _safe_csv_cell(-1000) == "-1000"


def _finance_records(db):
    student = Student(
        admission_number="LI-2026-00991",
        full_name="Export Student",
        mobile="9000000991",
        status="active",
    )
    db.add(student)
    db.flush()
    enrollment = Enrollment(
        student_id=student.id,
        program="JEE",
        batch="Tatva",
        enrollment_date=date(2026, 7, 1),
        status="active",
        is_active=True,
    )
    db.add(enrollment)
    db.flush()
    agreement = FeeAgreement(
        student_id=student.id,
        enrollment_id=enrollment.id,
        agreed_amount=80_000,
        legacy_registration_total=0,
        currency="INR",
        status="active",
    )
    db.add(agreement)
    db.flush()
    db.add(
        PaymentTransaction(
            student_id=student.id,
            fee_agreement_id=agreement.id,
            transaction_date=date(2026, 7, 2),
            amount=5_000,
            method="upi",
            transaction_type="payment",
            source_note="ERP receipt",
            status="posted",
            reconciliation_status="ready",
        )
    )
    db.commit()


def test_owner_can_export_student_and_fee_csv(
    client,
    database,
    owner_headers,
):
    _finance_records(database)
    students = client.get(
        "/api/reports/export/students",
        headers=owner_headers,
    )
    assert students.status_code == 200
    assert students.headers["content-type"].startswith("text/csv")
    assert "lakshya-students-" in students.headers["content-disposition"]
    assert "LI-2026-00991,Export Student" in students.text
    assert ",JEE,Tatva,2026-07-01,active," in students.text

    fees = client.get("/api/reports/export/fees", headers=owner_headers)
    assert fees.status_code == 200
    assert "80000,5000,75000,INR,active" in fees.text


def test_client_balance_reconciliation_changes_balance_not_cash_received(
    client,
    database,
    owner_headers,
):
    _finance_records(database)
    agreement = database.query(FeeAgreement).one()
    student = database.query(Student).filter_by(full_name="Export Student").one()
    database.add(
        PaymentTransaction(
            student_id=student.id,
            fee_agreement_id=agreement.id,
            legacy_import_id="client-snapshot-test",
            legacy_line_number=1,
            transaction_date=date(2026, 8, 3),
            amount=7_000,
            method="client_statement",
            transaction_type="balance_credit",
            source_note="Client-confirmed balance",
            status="posted",
            reconciliation_status="ready",
        )
    )
    database.commit()

    fees = client.get("/api/reports/export/fees", headers=owner_headers)
    assert fees.status_code == 200
    assert "80000,5000,68000,INR,active" in fees.text


def test_accounts_can_export_but_students_cannot(
    client,
    database,
    parent_headers,
):
    accounts = User(
        mobile="9000000998",
        full_name="Accounts User",
        role="accounts",
        password_hash=hash_password("Password123!"),
    )
    database.add(accounts)
    database.commit()
    accounts_headers = {
        "Authorization": f"Bearer {create_token(accounts)}",
    }
    assert client.get(
        "/api/reports/exports",
        headers=accounts_headers,
    ).status_code == 200
    assert client.get(
        "/api/reports/export/audit",
        headers=accounts_headers,
    ).status_code == 200
    assert client.get(
        "/api/reports/export/students",
        headers=parent_headers,
    ).status_code == 403


def test_report_date_range_validation(client, owner_headers):
    response = client.get(
        "/api/reports/export/attendance?from=2026-07-20&to=2026-07-01",
        headers=owner_headers,
    )
    assert response.status_code == 422


def test_report_overview_extends_confirmed_baseline_with_new_daily_register(
    client,
    database,
    owner_headers,
):
    student = Student(
        admission_number="LI-2026-00992",
        full_name="Attendance Report Student",
        mobile="9000000992",
        status="active",
    )
    database.add(student)
    database.flush()
    database.add(
        AttendancePeriodSummary(
            student_id=student.id,
            source_student_code="T-99",
            batch_name="Tatva",
            period_start=date(2026, 7, 1),
            period_end=date(2026, 8, 3),
            present_days=2,
            absent_days=1,
            working_days=3,
            attendance_rate=66.67,
            source_name="Client workbook",
            source_reference="test.xlsx",
            status="confirmed",
        )
    )
    owner = database.query(User).filter_by(role="owner").one()
    register = AttendanceRegister(
        register_kind="manual",
        attendance_date=date(2026, 8, 18),
        batch_name="Tatva",
        stream_name="__all__",
        subject_name="Daily attendance",
        status="submitted",
        submitted_by=owner.id,
    )
    database.add(register)
    database.flush()
    database.add(
        AttendanceEntry(
            register_id=register.id,
            student_id=student.id,
            status="present",
            reason="Signed paper register",
            marked_by=owner.id,
        )
    )
    database.commit()

    response = client.get("/api/reports/overview", headers=owner_headers)

    assert response.status_code == 200
    payload = response.json()
    assert payload["metrics"]["attendanceRate"] == 75.0
    assert {row["status"]: row["count"] for row in payload["attendance"]} == {
        "absent": 1,
        "present": 3,
    }
    workbook = _workbook(client, owner_headers, "attendance")
    assert workbook["Daily attendance"]["A7"].value == datetime(2026, 8, 18)
    summary = workbook["Period summaries"]
    assert summary["A7"].value == datetime(2026, 7, 1)
    assert summary["I7"].value == pytest.approx(0.6667)
    assert summary["I7"].number_format == "0.0%"


def test_populated_report_downloads_for_visual_review(client, database, owner_headers, tmp_path):
    """Keep synthetic exported files in pytest's temporary directory for visual QA."""
    _finance_records(database)
    owner = database.query(User).filter_by(role="owner").one()
    student = database.query(Student).one()
    batch = Batch(name="Tatva", program="JEE")
    subject = Subject(name="Physics", code="QA-PHY", program="JEE")
    room = Room(name="Test classroom")
    database.add_all([batch, subject, room])
    database.flush()
    session = ClassSession(batch_id=batch.id, subject_id=subject.id, faculty_id=owner.id, room_id=room.id,
        starts_at=datetime(2026, 9, 16, 5, tzinfo=timezone.utc), ends_at=datetime(2026, 9, 16, 6, tzinfo=timezone.utc))
    database.add(session)
    database.flush()
    class_register = AttendanceRegister(class_session_id=session.id, status="submitted")
    database.add(class_register)
    database.flush()
    database.add(AttendanceEntry(register_id=class_register.id, student_id=student.id, status="present", reason="Class register", marked_by=owner.id))
    register = AttendanceRegister(register_kind="manual", attendance_date=date(2026, 9, 16), batch_name="Tatva", stream_name="__all__", subject_name="Daily attendance", status="submitted")
    database.add(register)
    database.flush()
    database.add(AttendanceEntry(register_id=register.id, student_id=student.id, status="present", reason="Signed register", marked_by=owner.id))
    database.add(AttendancePeriodSummary(student_id=student.id, source_student_code="T-99", batch_name="Tatva", period_start=date(2026, 7, 1), period_end=date(2026, 8, 3), present_days=20, absent_days=4, working_days=24, attendance_rate=83.33, source_name="Client workbook", source_reference="sample.xlsx", status="confirmed"))
    database.add(AuditLog(actor_id=owner.id, action="student.updated", entity_type="student", entity_id=student.id, created_at=datetime(2026, 9, 16, 6, tzinfo=timezone.utc)))
    database.commit()
    for name in ("students", "fees", "attendance", "audit"):
        response = client.get(f"/api/reports/export/{name}?format=xlsx", headers=owner_headers)
        assert response.status_code == 200
        (tmp_path / f"{name}.xlsx").write_bytes(response.content)
    assert len(list(tmp_path.glob("*.xlsx"))) == 4
    sheet = load_workbook(tmp_path / "attendance.xlsx")["Class attendance"]
    assert sheet["D7"].value == "Tatva"
    assert sheet["E7"].value == "JEE"
    assert sheet["F7"].value == "Physics"
    assert sheet["J7"].value == datetime(2026, 9, 16, 10, 30)
